"""Knowledge ingestion: turn SOPs / business logic / product data into memory.

Answers "I don't know where the data is": upload files (pdf, docx, md, txt,
csv, json) or paste text on the Knowledge page. Each document is split into
chunks, every chunk is classified to a department agent (so the rulebook is
divided per agent), embedded into the vector index, and its concepts are added
to the knowledge graph (docs/architecture.md §2.6).
"""

import io
import json
import re

from sqlmodel import Session, select

from app.database import engine
from app.domain import DEPARTMENT_KEYWORDS
from app.memory import kg
from app.models import KnowledgeDocument, MemoryChunk
from app.utils import iso, utcnow

MAX_WORDS = 110
MIN_WORDS = 25


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(data))
        lines = []
        for p in d.paragraphs:
            if not p.text.strip():
                lines.append("")
            elif p.style and p.style.name.lower().startswith("heading"):
                lines.append("## " + p.text.strip())
            else:
                lines.append(p.text.strip())
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    text = data.decode("utf-8-sig", errors="replace")
    if name.endswith(".csv"):
        return _csv_to_text(text)
    if name.endswith(".json"):
        try:
            return _flatten_json(json.loads(text))
        except Exception:
            return text
    return text


def _csv_to_text(text: str) -> str:
    """A spreadsheet row is a fact: keep it whole, as its own paragraph ("column: value, column: value")."""
    import csv

    rows = list(csv.reader(io.StringIO(text)))
    if len(rows) < 2:
        return text
    head = [h.strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        cells = [f"{h}: {v.strip()}" for h, v in zip(head, r) if v.strip()]
        if cells:
            out.append(", ".join(cells))
    return "\n\n".join(out)


def _flatten_json(obj, prefix: str = "") -> str:
    if isinstance(obj, dict):
        return "\n".join(_flatten_json(v, f"{prefix}{k}: ") for k, v in obj.items())
    if isinstance(obj, list):
        return "\n\n".join(_flatten_json(v, prefix) for v in obj)
    return f"{prefix}{obj}"


def classify_department(text: str, default: str = "Other") -> str:
    low = text.lower()
    scores = {}
    for dept, words in DEPARTMENT_KEYWORDS.items():
        scores[dept] = sum(len(re.findall(rf"\b{re.escape(w)}", low)) for w in words)
    best, top = max(scores.items(), key=lambda kv: kv[1])
    return best if top >= 2 else default


def guess_kind(title: str, text: str) -> str:
    low = (title + " " + text[:600]).lower()
    if any(w in low for w in ("price", "specification", "sku", "warranty", "model", "catalog", "product")):
        return "product_db"
    if any(w in low for w in ("procedure", "steps", "sop", "how to", "troubleshoot", "policy")):
        return "sop"
    return "business_logic"


def chunk_text(title: str, text: str) -> list[tuple[str, str]]:
    """Split into (chunk_title, chunk_text) by headings, then by size."""
    heading = title
    blocks: list[tuple[str, str]] = []
    buf: list[str] = []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            blocks.append((heading, body))
        buf.clear()

    for line in text.replace("\r\n", "\n").split("\n"):
        m = re.match(r"^\s{0,3}#{1,4}\s+(.*)$", line)
        if m:
            flush()
            heading = f"{title} — {m.group(1).strip()}" if m.group(1).strip() != title else title
        else:
            buf.append(line)
    flush()

    chunks: list[tuple[str, str]] = []
    for h, body in blocks:
        paras = []
        for p_ in (x.strip() for x in re.split(r"\n\s*\n", body) if x.strip()):
            if len(p_.split()) > MAX_WORDS * 1.4:
                sentences = re.split(r"(?<=[.!?])\s+", p_.replace("\n", " "))
                paras.extend(s_ for s_ in sentences if s_)
            else:
                paras.append(p_)
        cur: list[str] = []
        words = 0
        for p in paras:
            w = len(p.split())
            if cur and words + w > MAX_WORDS:
                chunks.append((h, "\n".join(cur)))
                cur, words = [], 0
            cur.append(p)
            words += w
        if cur:
            chunks.append((h, "\n".join(cur)))

    # merge tiny neighbours that share a heading
    merged: list[tuple[str, str]] = []
    for h, c in chunks:
        if merged and merged[-1][0] == h and len(merged[-1][1].split()) < MIN_WORDS:
            merged[-1] = (h, merged[-1][1] + "\n" + c)
        else:
            merged.append((h, c))
    return merged


def ingest_text(
    title: str,
    text: str,
    *,
    kind: str | None = None,
    department: str | None = None,
    store: str = "common",
    source_name: str | None = None,
    category: str | None = None,
) -> dict:
    from app.memory.memory_engine import memory  # local import: memory_engine imports this module

    text = text.strip()
    if not text:
        raise ValueError("document is empty")
    kind = kind or guess_kind(title, text)
    # Product data is shared knowledge every agent may use, so it lives under "Other" unless told otherwise.
    doc_default = department or ("Other" if kind == "product_db" else classify_department(title + "\n" + text))
    pieces = chunk_text(title, text)
    if not pieces:
        raise ValueError("no readable content found in the document")

    per_dept: dict[str, int] = {}
    with Session(engine) as s:
        doc = KnowledgeDocument(
            title=title, kind=kind, category=category or _category(kind, doc_default),
            department=doc_default, source_name=source_name or title,
        )
        s.add(doc)
        s.commit()
        s.refresh(doc)
        for ctitle, ctext in pieces:
            dept = department or (doc_default if kind == "product_db" else classify_department(ctitle + "\n" + ctext, default=doc_default))
            s.add(MemoryChunk(store=store, kind=kind, department=dept, title=ctitle, text=ctext,
                              document_id=doc.document_id, source=title))
            per_dept[dept] = per_dept.get(dept, 0) + 1
            kg.index_chunk(s, doc_label=title, department=dept, text=ctitle + " " + ctext, store=store)
        doc.chunk_count = len(pieces)
        doc.last_updated = utcnow()
        # documents are mostly owned by the department holding most chunks
        doc.department = max(per_dept.items(), key=lambda kv: kv[1])[0]
        doc.used_by_agent = doc.department
        s.add(doc)
        s.commit()
        result = {"document_id": doc.document_id, "title": title, "kind": kind, "store": store,
                  "chunks": len(pieces), "departments": per_dept}
    memory.reload()
    return result


def _category(kind: str, dept: str) -> str:
    if kind == "product_db":
        return "Product Information"
    return {"Billing": "Billing", "Order": "Orders", "Technical": "Troubleshooting"}.get(dept, "Policies")


def ingest_file(filename: str, data: bytes, **kw) -> dict:
    title = re.sub(r"\.[A-Za-z0-9]+$", "", filename).replace("_", " ").replace("-", " ").strip().title()
    return ingest_text(title, extract_text(filename, data), source_name=filename, **kw)


def list_documents() -> list[dict]:
    with Session(engine) as s:
        docs = s.exec(select(KnowledgeDocument).order_by(KnowledgeDocument.document_id)).all()
    return [
        {
            "document_id": d.document_id, "title": d.title, "category": d.category, "kind": d.kind,
            "department": d.department, "chunk_count": d.chunk_count,
            "referenced_in_tickets": d.referenced_in_tickets, "source_name": d.source_name,
            "last_updated": iso(d.last_updated),
        }
        for d in docs
    ]


def delete_document(document_id: int) -> bool:
    from app.memory.memory_engine import memory

    with Session(engine) as s:
        doc = s.get(KnowledgeDocument, document_id)
        if not doc:
            return False
        for c in s.exec(select(MemoryChunk).where(MemoryChunk.document_id == document_id)).all():
            s.delete(c)
        s.delete(doc)
        s.commit()
    memory.reload()
    return True


